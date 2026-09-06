"""
04_codebook.py -- SAS: generate CODEBOOK_SAS.xlsx, the dataset's only codebook, in the Dropbox dataset folder next to README.txt.

Per variable: label, storage type, value labels (first few), non-missing count per wave, the
source per wave recorded by 01_clean (person/household files), version decisions and label
variants recorded by 02_merge.
"""
import json, re, time
import pandas as pd
from sas_helpers import paths, get_logger, read_dta, read_meta, LOGS, HERE, DATASET

TAG = "SAS"
log = get_logger("04_codebook"); P = paths()
align = json.load(open(LOGS / "merge_alignment.json"))
metas = {f.stem.replace("clean_", "").replace("_meta", ""): json.load(open(f)) for f in LOGS.glob("clean_*_meta.json")}   # year -> {"S/stem": {source, rows, ...}}

def chunked_counts(path, key, chunk=250_000):
    """non-missing counts per column by `key` (wave/year), reading the file in row chunks so
    multi-GB pooled files never sit in memory whole; returns (counts DataFrame, ordered key values, n rows)."""
    m = read_meta(path); counts, order = None, []
    for off in range(0, m.number_rows, chunk):
        part, _, _ = read_dta(path, row_offset=off, row_limit=chunk)
        if key not in part.columns: part[key] = "all"
        c = part.groupby(key, sort=False).count(); c[key] = part.groupby(key, sort=False).size()   # the key itself: group size
        counts = c if counts is None else counts.add(c, fill_value=0)
        order += [k for k in part[key].unique().tolist() if k not in order]
    return counts.reindex(order).fillna(0).astype(int), order, int(m.number_rows)


FILE_UNIVERSE = {
    "SAS_pooled_plotcrop.dta": "2017-2025: every crop-PRODUCTION record of the plot questionnaire (part II) in sampled segments (small-scale, area frame) and on large-scale farms (list frame, farm_type); seasons A (Sep-Feb), B (Mar-Aug), C (marshland). Crop codes are the wave's own list (crop_list: 2017-2019 vs 2020+; text in crop_name). (wave, farm_type, segment, holder, plot, crop) is NOT unique (a few dozen duplicated rows per wave, as shipped). The official cultivated-area universe is the screening crop record (2_Intermediate/appended/SAS_pooled_screening_crops.dta), which reproduces the published maize areas 2020-2023",
    "SAS_pooled_screening_crops.dta": "2019-2025: every screened plot x crop record (planted crops) of the screening questionnaire in sampled segments and on large-scale farms -- the CULTIVATED-AREA universe: sum(wt x crop_area) by crop reproduces the published seasonal areas 2020 A - 2023 B (checks, section D); crop = the wave's own list (crop_list, crop_name); 2019 = the single screening file, 2020-2023 = screening_crops, 2024-2025 = screening",
    "SAS_pooled_plotcrop_2013_2016.dta": "2013-2016: plot x crop records of every shipped file with a plot-level crop record: screening, crop-area, sowing / production / harvest and plot-roster files (record_type, source_module); several record types per wave -> keys are unique only within a record type, if at all; 2013 keys are tract and plot within segment; crop = the wave's own list (crop_list)",
}
FILE_UNIVERSE_DEFAULT = 'plots of sampled segments (area frame, strata 10 hillside / 20 marshland / 30 rangeland / 40 mixed) and of large-scale farmers (list frame, stratum 90); one record per plot x crop x season'

def sheet_name(fname, used):
    """Excel sheet name for a file: the stem without the <TAG>_pooled_ prefix, <= 31 chars, unique."""
    s = fname[:-4] if fname.endswith(".dta") else fname
    for pre in (f"{TAG}_pooled_", f"{TAG}_"):
        if s.startswith(pre): s = s[len(pre):]; break
    s = re.sub(r"[\[\]:*?/\\]", "_", s)[:31] or "sheet"
    base, i = s, 2
    while s in used: s = f"{base[:28]}_{i}"; i += 1
    used.add(s); return s


def write_workbook(path, notes, books):
    """One workbook per dataset (its only codebook, next to README.txt on Dropbox): README (notes), files
    (overview), ONE SHEET PER FINAL DATASET (variable table), modules (every appended module file of
    2_Intermediate/appended/ stacked in one sheet, file in the first column), value_labels (every code of
    every labelled variable, all files) and checks (the latest logs/checks_report.txt)."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    used = {"README", "files", "modules", "value_labels", "checks"}
    finals = [(cb, m) for cb, m, _ in books if m is not None and m["folder"] == "3_Final"]
    mods = [(cb, m) for cb, m, _ in books if m is not None and m["folder"] != "3_Final"]
    names = {m["file"]: sheet_name(m["file"], used) for _, m in finals}
    names.update({m["file"]: "modules" for _, m in mods})
    files = pd.DataFrame([dict(m, sheet=names[m["file"]]) for _, m, _ in books if m is not None])
    vls = pd.DataFrame([r for _, _, v in books for r in v], columns=["file", "variable", "code", "label"])
    report = LOGS / "checks_report.txt"
    checks = pd.DataFrame({"checks_report": report.read_text().splitlines() if report.exists() else ["(no logs/checks_report.txt yet: run 03_checks.py)"]})
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        pd.DataFrame({"README": notes}).to_excel(xw, sheet_name="README", index=False)
        files.to_excel(xw, sheet_name="files", index=False)
        for cb, m in finals: cb.to_excel(xw, sheet_name=names[m["file"]], index=False)
        if mods:
            stack = pd.concat([cb.assign(file=m["file"]) for cb, m in mods], ignore_index=True, sort=False)
            lead = [c for c in ("file", "variable", "label", "type", "value_labels", "n_value_labels", "universe") if c in stack.columns]
            ncols = sorted(c for c in stack.columns if c.startswith("n_") and c not in lead)
            stack = stack[lead + ncols + [c for c in stack.columns if c not in lead + ncols]]
            stack.to_excel(xw, sheet_name="modules", index=False)
        vls.to_excel(xw, sheet_name="value_labels", index=False)
        checks.to_excel(xw, sheet_name="checks", index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "A2"
            for cell in ws[1]: cell.font = Font(bold=True)
            for j, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), start=1):
                w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
                ws.column_dimensions[get_column_letter(j)].width = min(max(10, w + 2), 80)
    log.info("wrote %s (%d final-dataset sheets, %d module files in one sheet, %d value-label rows)", path, len(finals), len(mods), len(vls))

def codebook(fname, info):
    path = P["root"] / info.get("dir", "3_Final") / fname          # unit files in 3_Final, modules in 2_Intermediate/appended
    df, vl, vv = read_dta(path, row_limit=1); m = read_meta(path); types = m.readstat_variable_types
    nn, waves, nrows = chunked_counts(path, "wave")                    # chunked: pooled files can be several GB
    decisions = info.get("decisions", {}); rows = []
    unit = info.get("unit", "")
    for c in df.columns:
        base = next((b for b, d in decisions.items() if c in d.get("versions", {})), c)
        d = decisions.get(base, {}); vers = d.get("versions", {}); ws_here = vers.get(c, [])
        labs = {w: l for w, l in (d.get("labels_by_wave") or {}).items() if w in ws_here}
        variants = "; ".join(f"{w}: {l[:50]}" for w, l in labs.items() if l and l != d.get("reference_label")) if len(set(labs.values())) > 1 else ""
        others = ", ".join(f"{n} ({', '.join(g)})" for n, g in vers.items() if n != c) if len(vers) > 1 else ""
        src = []
        for w in (ws_here or info.get("waves", [])):                     # lineage: wave -> (year, "S/stem") -> 01_clean's source map
            y, mk = info.get("wave_modules", {}).get(w, (None, None))
            sm = metas.get(y, {}).get(mk, {}).get("source", {}) if y else {}
            if c in sm or base in sm: src.append(f"{w}: {sm.get(c, sm.get(base))}")
        src = "; ".join(src)
        vlab = vv.get(c, {}); vtxt = ", ".join(f"{k}={v}" for k, v in list(vlab.items())[:5]) + (" …" if len(vlab) > 5 else "")
        conf = info.get("value_label_conflicts", {}).get(c)
        rows.append({"variable": c, "label": vl.get(c) or "", "type": types.get(c, ""), "value_labels": vtxt, "n_value_labels": len(vlab),
                     **{f"n_{w}": int(nn.loc[w, c]) if w in nn.index and c in nn.columns else 0 for w in waves},
                     "source": src, "versions": others, "label_variants": variants, "value_label_conflicts": json.dumps(conf) if conf else ""})
    vl_rows = [{"file": fname, "variable": c, "code": k, "label": v} for c in df.columns for k, v in vv.get(c, {}).items()]
    return pd.DataFrame(rows), {"file": fname, "folder": info.get("dir", "3_Final"), "unit": unit, "rows": nrows, "variables": df.shape[1],
                                "waves": ", ".join(map(str, waves)), "universe": FILE_UNIVERSE.get(fname, FILE_UNIVERSE_DEFAULT)}, vl_rows

NOTES = [f"SAS codebook -- generated by rwa-trs/NISR/{DATASET}/04_codebook.py on {time.strftime('%Y-%m-%d %H:%M')} from 3_Final/ (and 2_Intermediate/appended/). Do not hand-edit: re-run python master.py 04.",
         "Sheets: files = one row per final/appended file (folder, unit, rows, variables, waves, universe, sheet); one sheet per FINAL dataset = one row per variable (label, storage type, first value labels, number of value labels, non-missing count per wave, source variable per wave, versions, label variants, value-label text conflicts); modules = the same table for every appended module file of 2_Intermediate/appended/, stacked, file in the first column; value_labels = every code of every labelled variable (all files); checks = the latest verification report (logs/checks_report.txt).",
         'Key block: survey year season wave farm_type prov dist stratum segment holder plot crop wt, plus the CORE quantities plot_area_ha crop_area_ha harvested_area_ha production_kg yield_kg_ha (explicit per-year map in 02_merge.py, recorded in logs/merge_alignment.json).',
         f"Alignment rule: a same-named variable is one column across waves only when (a) its variable labels are similar (token Jaccard >= {align['threshold']}) AND (b) its value labels are compatible (no code whose text means something else, similarity >= {align.get('vl_threshold', 0.6)}; an unlabelled wave's values inside the labelled range); the largest group keeps the name, the others are <name>_v2, _v3 ... (split reasons in logs/merge_alignment.json); forced splits by year: {', '.join(f'{k} {v}' for k, v in (align.get('force_split') or {}).items()) or 'none'}. The crop code is exempt: it keeps the wave's own list (crop_list, crop_name) and no pooled value labels.",
         "Value labels are the union over waves within a version; where the same code had different text the most recent wave's text is kept and the conflict listed.",
         'universe (files sheet): from the questionnaires; no sampling/weighting document exists in z_Documentation (strata 10 hillside / 20 marshland / 30 rangeland / 40 mixed / 90 large-scale list frame).',
         f"Processing notes (what z_Documentation says and how it was applied; every decision and why) are kept in the project memory file NISR-{DATASET}.md (Green Jobs - TRS folder), not in git or Dropbox."]
books = [codebook(fname, info) for fname, info in align["files"].items()]
# the crop code keeps no pooled value labels (the lists differ by wave): every wave's own crop dictionary goes into the value_labels sheet, file = the cleaned per-wave file
for y, m in sorted(metas.items()):
    for mk, i in sorted(m.items()):
        d = i.get("value_labels", {}).get("crop", {})
        if len(d) >= 20: books.append((None, None, [{"file": i["out"], "variable": "crop", "code": k, "label": v} for k, v in d.items()]))
write_workbook(P["root"] / f"CODEBOOK_{TAG}.xlsx", NOTES, books)

# ---------------------------------------------------------------- README.txt on Dropbox
# The dataset folder's README.txt keeps its hand-written sections; the block between the two
# marker lines below is regenerated from 3_Final/ and 2_Intermediate/ on every run, and the
# pre-pipeline sentences that are now stale are replaced by a pointer to the git folder.
def update_readme():
    import re, time
    from pathlib import Path
    readme = P["root"] / "README.txt"
    if not readme.exists(): log.warning("no README.txt in %s; skipped", P["root"]); return
    start, end = "=== PROCESSED OUTPUTS (generated by 04_codebook.py; do not edit inside) ===", "=== END PROCESSED OUTPUTS ==="
    finals = sorted(P["final"].glob("*.dta")); inters = sorted(P["inter"].glob("*.dta"))
    lines = [start, f"Pipeline: rwa-trs/NISR/{DATASET}/  (python master.py)   last run: {time.strftime('%Y-%m-%d %H:%M')}",
             f"Codebook: CODEBOOK_{TAG}.xlsx in this folder (sheets: README, files, one per final dataset, modules, value_labels, checks).",
             f"Code and run logs: git repository rwa-trs/NISR/{DATASET}/. Processing notes (what z_Documentation says and how it was applied; every decision and why): project memory file NISR-{DATASET}.md (Green Jobs - TRS folder, outside git and Dropbox).", "",
             f"3_Final/  ({len(finals)} files)"]
    for f in finals:
        try:
            m = read_meta(f); lines.append(f"  {f.name:55s} {m.number_rows:>10,} rows x {len(m.column_names):>4} vars   {m.file_label or ''}")
        except Exception as e: lines.append(f"  {f.name:55s} (unreadable: {e})")
    app = sorted((P["inter"] / "appended").glob("*.dta")) if (P["inter"] / "appended").exists() else []
    lines += ["", f"2_Intermediate/  ({len(inters)} files: one cleaned file per wave and unit/module; see CODEBOOK_{TAG}.xlsx)"]
    if app:
        lines.append(f"2_Intermediate/appended/  ({len(app)} module-level files appended across waves)")
        for f in app:
            try: m = read_meta(f); lines.append(f"  {f.name:55s} {m.number_rows:>10,} rows x {len(m.column_names):>4} vars")
            except Exception as e: lines.append(f"  {f.name:55s} (unreadable: {e})")
    lines.append(end)
    txt = readme.read_text(encoding="utf-8", errors="replace")
    for stale in ["2_Intermediate/ and 3_Final/ are empty: no clean.do / merge.do written for\nthis survey yet. The Census, EICV and LFS folders have working examples.",
                  "2_Intermediate/  0 files (not built yet)", "3_Final/         0 files (not built yet)"]:
        txt = txt.replace(stale, "Processed by the Python pipeline in rwa-trs/NISR/ (see the PROCESSED OUTPUTS block).")
    txt = re.sub(r"Code lives in code/Publicly Available Data - NISR/ and was repointed to these\npaths on the same date \(previous paths under \"Rwanda - Tourism, Green Jobs/\nData/Public\" were already stale\).",
                 "Code now lives in the git repository rwa-trs/NISR/<dataset>/ (Python); the Stata\ndo-files are archived in Rwanda - TRS/_archive/code-NISR-stata-2026-09-04/.", txt)
    new_block = "\n".join(lines)
    if start in txt and end in txt:
        txt = txt[:txt.index(start)] + new_block + txt[txt.index(end) + len(end):]
    else:
        head_end = txt.find("\n\n") + 2 if "\n\n" in txt[:600] else 0
        txt = txt[:head_end] + new_block + "\n\n" + txt[head_end:]
    readme.write_text(txt, encoding="utf-8"); log.info("README.txt updated (%s)", readme)

update_readme()
